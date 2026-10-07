"""Patch the user's own edited AUTOBONO_Database workbook (their column order is kept).

1. Brings the notes (cell comments) of the old AUTOBONO sheet across: cost breakdowns go into the
   Expenses descriptions and onto the Cars cost cells; car / buyer / sale-price notes onto the
   matching Cars cells.
2. "Sold" now means "Sale date is filled", so a car marked RESERVED (sold, not finished yet)
   counts as sold in every total while keeping its yellow status.

Run:  python3 update_user_workbook.py OLD_AUTOBONO.xlsx USER_DATABASE.xlsx OUT.xlsx
"""
import re
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment

OLD, USER, OUT = sys.argv[1:4]
old = openpyxl.load_workbook(OLD)
wb = openpyxl.load_workbook(USER)
cars = wb['Cars']
H = {c.value: c.column for c in cars[1] if c.value}
row_of = {int(cars.cell(r, 1).value): r for r in range(2, cars.max_row + 1)
          if isinstance(cars.cell(r, 1).value, (int, float))}


def clean(text):
    """Drop Excel's threaded-comment boilerplate, keep the user's words."""
    if '[Threaded comment]' in text:
        text = text.split('Comment:', 1)[-1]
        text = re.sub(r'\n\s*Reply:\s*\n', '\n', text)
    lines = [ln.strip() for ln in text.replace('\t', ' ').splitlines()]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return '\n'.join(lines)


def note(ws, coord):
    c = ws[coord]
    return clean(c.comment.text) if c.comment and clean(c.comment.text) else None


# old column → new Cars header
MAP_2024 = {'D': 'Model', 'J': 'VIN', 'L': 'To do', 'M': 'Buyer', 'N': 'Country (old "State")',
            'V': 'Sale price (€)', 'X': 'Profit (€)', 'Z': 'Investor / partner share (€)', 'AF': 'Online since'}
MAP_2023 = {'C': 'Model', 'J': 'Buyer', 'R': 'Sale price (€)'}
COST_HEADERS = ('Purchase price (€)', 'Extra costs (€) [from Expenses]')

notes = {}       # car id → {header: text}
cost_note = {}   # car id → text


def add(no, header, text):
    if text:
        cur = notes.setdefault(no, {}).get(header)
        notes[no][header] = text if not cur or text in cur else f'{cur}\n—\n{text}'


ws24 = old['2024']
in_2024 = set()
for r in range(5, 173):
    no = ws24.cell(r, 3).value
    if not isinstance(no, (int, float)) or not ws24.cell(r, 4).value:
        continue
    no = int(no)
    in_2024.add(no)
    for col, header in MAP_2024.items():
        add(no, header, note(ws24, f'{col}{r}'))
    texts = [note(ws24, f'R{r}')]
    link = re.fullmatch(r'=\$?R\$?(\d+)', str(ws24[f'R{r}'].value or '').replace(' ', ''))
    if link:                                    # "=R190": breakdown sits on the old 2023 copy below
        texts.append(note(ws24, f'R{link.group(1)}'))
    texts = [t for t in texts if t]
    if texts:
        cost_note[no] = '\n—\n'.join(dict.fromkeys(texts))

ws23 = old['2023']
for r in range(4, 20):
    no = ws23.cell(r, 2).value
    if not isinstance(no, (int, float)) or not ws23.cell(r, 3).value:
        continue
    no = int(no)
    if no in in_2024:                           # continued on the 2024 tab — its notes win
        continue
    for col, header in MAP_2023.items():
        add(no, header, note(ws23, f'{col}{r}'))
    if note(ws23, f'N{r}'):
        cost_note[no] = note(ws23, f'N{r}')


def put(cell, text):
    c = Comment(text, 'AUTOBONO (old sheet)')
    c.width, c.height = 320, min(60 + 15 * text.count('\n'), 420)
    cell.comment = c


placed = 0
for no, by_header in notes.items():
    if no not in row_of:
        continue
    for header, text in by_header.items():
        if header in H:
            put(cars.cell(row_of[no], H[header]), text)
            placed += 1
for no, text in cost_note.items():
    if no in row_of:
        for header in COST_HEADERS:
            put(cars.cell(row_of[no], H[header]), f'Costs as noted in the old sheet:\n{text}')
            placed += 1

# Expenses: replace the "Imported: …" description with the real breakdown
ex = wb['Expenses']
described = 0
for r in range(2, ex.max_row + 1):
    car, desc = ex.cell(r, 2).value, ex.cell(r, 4).value
    if not (isinstance(car, (int, float)) and isinstance(desc, str) and desc.startswith('Imported:')):
        continue
    formula = desc.split('old cost formula ', 1)[-1]
    text = cost_note.get(int(car))
    if text:
        ex.cell(r, 4).value = f'{text}\n(old cost formula: {formula})'
        ex.cell(r, 4).alignment = Alignment(wrap_text=True, vertical='top')
        described += 1
    else:
        ex.cell(r, 4).value = f'Extra costs (no description in the old sheet) — old cost formula: {formula}'
ex.column_dimensions['D'].width = max(ex.column_dimensions['D'].width or 0, 70)

# "Sold" = Sale date filled, so RESERVED-with-sale-date counts as sold
status = f'Cars!${cars.cell(1, H["Status"]).column_letter}$2:${cars.cell(1, H["Status"]).column_letter}$1000'
sdate_col = cars.cell(1, H['Sale date']).column_letter
sdate = f'Cars!${sdate_col}$2:${sdate_col}$1000'
st_col = cars.cell(1, H['Status']).column_letter
changed = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not (isinstance(v, str) and v.startswith('=') and 'SOLD' in v):
                continue
            nv = (v.replace(f'{status},"<>SOLD"', f'{sdate},""')
                   .replace(f'{status},"SOLD"', f'{sdate},"<>"'))
            nv = re.sub(rf'IF\({st_col}(\d+)<>"SOLD",""', rf'IF({sdate_col}\1="",""', nv)
            if nv != v:
                c.value = nv
                changed += 1
for cf in cars.conditional_formatting:
    for rule in cf.rules:
        rule.formula = [f.replace(f'${st_col}2<>"SOLD"', f'${sdate_col}2=""') for f in rule.formula]

# wording on Summary / Read me
sm = wb['Summary']
for row in sm.iter_rows():
    for c in row:
        if c.value == 'Cars reserved':
            c.value = 'RESERVED (sold or promised, not finished yet)'
rm = wb['Read me']
for row in rm.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith('• Car sold: Status = SOLD'):
            c.value = ('• Car sold: fill Sale date, Sale price, Sold via, "VAT on sale" and Buyer — from that moment it counts as sold '
                       'everywhere. Status = SOLD when everything is dealt with (row turns grey); keep it RESERVED (yellow) while '
                       'something is still open.')
        elif isinstance(c.value, str) and c.value.startswith('• Every extra cost'):
            c.value = ('• Every extra cost is a line on the Expenses tab with the Car ID and what it was for. Old cars: the description '
                       'is your old note on the cost cell; hover over Purchase price / Extra costs on Cars to see it too.')

# ---------------------------------------------------------------- future-proofing
import copy
import datetime as dt
from openpyxl.formula.translate import Translator

YEAR_ROWS, BRAND_ROWS = 15, 40          # room for years 2023–2037 and 40 brands
C = lambda h: f'Cars!${cars.cell(1, H[h]).column_letter}$2:${cars.cell(1, H[h]).column_letter}$1000'  # noqa: E731
SOLD = f'{C("Sale date")},"<>"'
UNSOLD = f'{C("Sale date")},"",{C("Status")},"<>IN USE",{C("Car ID")},"<>"'

# Lists: brands in order of importance, then the rest alphabetically (new brands: add at the bottom)
lists = wb['Lists']
TOP = ['VW', 'Škoda', 'Seat', 'Audi', 'Cupra', 'Mercedes-Benz', 'Mazda', 'Alfa Romeo']
brands = [lists.cell(r, 3).value for r in range(2, 61) if lists.cell(r, 3).value]
brands = [b for b in TOP if b in brands or b == 'Cupra'] + sorted(b for b in brands if b not in TOP)
for r in range(2, 61):
    lists.cell(r, 3).value = brands[r - 2] if r - 2 < len(brands) else None

# Summary: rebuild everything from "Sold cars by year" down, keeping the look
sm = wb['Summary']
st_title, st_head, st_label, st_num, st_total = (copy.copy(sm[c]._style) for c in ('B19', 'B20', 'B21', 'C21', 'B25'))
fmt = {k: sm[c].number_format for k, c in (('n', 'C21'), ('eur', 'D21'), ('days', 'F21'))}
for row in sm.iter_rows(min_row=19, max_row=max(sm.max_row, 120)):
    for c in row:
        c.value = None
        c._style = copy.copy(sm['A1']._style)


def cell(r, col, value, style, number_format=None):
    c = sm.cell(row=r, column=col, value=value)
    c._style = copy.copy(style)
    if number_format:
        c.number_format = number_format
    return c


r = 19
cell(r, 2, 'Sold cars by year (by sale date) — a new year appears by itself on 1 January', st_title)
r += 1
for i, h in enumerate(['Year', 'Cars sold', 'Profit (€)', 'Avg profit / car (€)', 'Avg days in stock', 'Avg days online',
                       'Investor / partner share (€)']):
    cell(r, 2 + i, h, st_head)
first = r + 1
for i in range(YEAR_ROWS):
    r += 1
    year = '2023' if i == 0 else f'=IF(B{r - 1}="","",IF(VALUE(B{r - 1})+1>YEAR(TODAY()),"",TEXT(VALUE(B{r - 1})+1,"0")))'
    cell(r, 2, year, st_label)
    crit = f'{SOLD},{C("Sale date")},">="&DATE(VALUE(B{r}),1,1),{C("Sale date")},"<"&DATE(VALUE(B{r})+1,1,1)'
    wrap = lambda f: f'=IF(B{r}="","",{f})'  # noqa: E731
    for col, f, nf in ((3, f'COUNTIFS({crit})', fmt['n']), (4, f'SUMIFS({C("Profit (€)")},{crit})', fmt['eur']),
                       (5, f'IF(C{r}=0,"",D{r}/C{r})', fmt['eur']),
                       (6, f'IFERROR(AVERAGEIFS({C("Days in stock (total)")},{crit}),"")', fmt['days']),
                       (7, f'IFERROR(AVERAGEIFS({C("Days online")},{crit}),"")', fmt['days']),
                       (8, f'SUMIFS({C("Investor / partner share (€)")},{crit})', fmt['eur'])):
        cell(r, col, wrap(f), st_num, nf)
r += 1
cell(r, 2, 'All years', st_total)
for col, f, nf in ((3, f'=SUM(C{first}:C{r - 1})', fmt['n']), (4, f'=SUM(D{first}:D{r - 1})', fmt['eur']),
                   (5, f'=IF(C{r}=0,"",D{r}/C{r})', fmt['eur']), (8, f'=SUM(H{first}:H{r - 1})', fmt['eur'])):
    cell(r, col, f, st_total, nf)

r += 2
cell(r, 2, 'Sold cars by brand (all years) — brands come from the Lists tab; add a new brand there', st_title)
r += 1
for i, h in enumerate(['Brand', 'Cars sold', 'Profit (€)', 'Avg profit / car (€)', 'Avg days in stock', 'In stock now']):
    cell(r, 2 + i, h, st_head)
bfirst = r + 1
for i in range(BRAND_ROWS):
    r += 1
    cell(r, 2, f'=IF(Lists!$C${2 + i}="","",Lists!$C${2 + i})', st_label)
    crit = f'{SOLD},{C("Brand")},B{r}'
    wrap = lambda f: f'=IF(B{r}="","",{f})'  # noqa: E731
    for col, f, nf in ((3, f'COUNTIFS({crit})', fmt['n']), (4, f'SUMIFS({C("Profit (€)")},{crit})', fmt['eur']),
                       (5, f'IF(C{r}=0,"",D{r}/C{r})', fmt['eur']),
                       (6, f'IFERROR(AVERAGEIFS({C("Days in stock (total)")},{crit}),"")', fmt['days']),
                       (7, f'COUNTIFS({UNSOLD},{C("Brand")},B{r})', fmt['n'])):
        cell(r, col, wrap(f), st_num, nf)
r += 1
cell(r, 2, 'Brand missing or not in Lists', st_label)
cell(r, 3, f'=COUNTIFS({SOLD})-SUM(C{bfirst}:C{r - 1})', st_num, fmt['n'])
cell(r, 4, f'=SUMIFS({C("Profit (€)")},{SOLD})-SUM(D{bfirst}:D{r - 1})', st_num, fmt['eur'])
sm.conditional_formatting.add(f'B{bfirst}:G{r - 1}', openpyxl.formatting.rule.FormulaRule(
    formula=[f'AND($B{bfirst}<>"",$C{bfirst}=0,$G{bfirst}=0)'], font=openpyxl.styles.Font(color='BFBFBF')))

r += 2
cell(r, 2, 'Expenses — who paid (to settle between partners)', st_title)
r += 1
cell(r, 2, 'Paid by', st_head)
cell(r, 3, 'Total (€)', st_head)
ex_rng = lambda col: f'Expenses!${col}$2:${col}$3000'  # noqa: E731
for i in range(6):
    r += 1
    cell(r, 2, f'=IF(Lists!$I${2 + i}="","",Lists!$I${2 + i})', st_label)
    cell(r, 3, f'=IF(B{r}="","",SUMIFS({ex_rng("E")},{ex_rng("F")},B{r},{ex_rng("B")},"<>EXAMPLE"))', st_num, fmt['eur'])
r += 1
cell(r, 2, '(not filled in)', st_label)
cell(r, 3, f'=SUMIFS({ex_rng("E")},{ex_rng("F")},"",{ex_rng("B")},"<>EXAMPLE",{ex_rng("E")},"<>")', st_num, fmt['eur'])

# Monthly: months up to December 2037
mo = wb['Monthly']
last = max(rr for rr in range(19, mo.max_row + 1) if isinstance(mo.cell(rr, 2).value, dt.datetime))
d = mo.cell(last, 2).value
rr = last
while (d.year, d.month) < (2037, 12):
    d = dt.datetime(d.year + (d.month == 12), d.month % 12 + 1, 1)
    rr += 1
    for col in range(2, 13):
        src = mo.cell(last, col)
        dst = mo.cell(rr, col)
        dst._style = copy.copy(src._style)
        dst.value = d if col == 2 else Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
for cf in mo.conditional_formatting:
    if 'B19' in str(cf.sqref):
        cf.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f'B19:L{rr}')

wb.save(OUT)
print(f'notes placed: {placed}, expense lines described: {described}, formulas switched to sale date: {changed}')
